#!/usr/bin/env python3
"""Parse verified Digital-ESD full text through the generic SLR scholarly parser.

Thin dashi_agda application wrapper.  Core parsing semantics stay in the SLR
repository.

The wrapper bridges the verified P0-G TSV to the generic scholarly parser while
preserving the crucial revision distinction:

  retrieved artifact bytes
      -> artifact SHA-256 / artifact revision
      -> text materialisation
      -> text SHA-256 / parser revision
      -> generic scholarly parser
      -> explicit parse receipt

Text/span anchors therefore belong to the materialised text revision rather
than being falsely attributed to raw PDF/DOCX bytes.

Outputs:
  materialization-receipts.jsonl
  scholarly-parser-input.jsonl
  slr-handoff-receipts.jsonl
  slr-parse-receipts.jsonl
  parser/{requests.jsonl,parser-output.jsonl,verified.jsonl}
  study_processing_census_with_parse.json

Authority:
  verified bytes != parsed study
  parsed study != reviewed evidence
  parsed study != SourceAuditAdmission
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html.parser
import json
import os
import subprocess
import sys
from pathlib import Path
from typing import Any


TEXT_EXTENSIONS = {".txt", ".md", ".tex", ".csv", ".tsv", ".json", ".jsonl"}
HTML_EXTENSIONS = {".html", ".htm"}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return [dict(row) for row in csv.DictReader(fh, delimiter="\t")]


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not path.exists():
        return rows
    with path.open("r", encoding="utf-8") as fh:
        for n, line in enumerate(fh, 1):
            if not line.strip():
                continue
            row = json.loads(line)
            if not isinstance(row, dict):
                raise ValueError(f"{path}:{n}: expected object")
            rows.append(row)
    return rows


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


class _HTMLTextExtractor(html.parser.HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript"}:
            self._skip_depth += 1
        elif tag.lower() in {"p", "br", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript"} and self._skip_depth:
            self._skip_depth -= 1
        elif tag.lower() in {"p", "div", "section", "article", "li", "tr"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and data.strip():
            self.parts.append(data)

    def text(self) -> str:
        lines = [" ".join(part.split()) for part in "".join(self.parts).splitlines()]
        return "\n".join(line for line in lines if line).strip() + "\n"


def extract_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(path))
        pages = []
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            pages.append(f"\n[PAGE {index}]\n{text}")
        return "".join(pages).strip() + "\n", "pypdf"
    except ImportError:
        pass

    try:
        import fitz  # type: ignore

        doc = fitz.open(str(path))
        pages = []
        for index, page in enumerate(doc, 1):
            pages.append(f"\n[PAGE {index}]\n{page.get_text('text')}")
        return "".join(pages).strip() + "\n", "pymupdf"
    except ImportError as exc:
        raise RuntimeError(
            "PDF text extraction requires pypdf or PyMuPDF/fitz"
        ) from exc


def extract_docx(path: Path) -> tuple[str, str]:
    try:
        from docx import Document  # type: ignore
    except ImportError as exc:
        raise RuntimeError("DOCX text extraction requires python-docx") from exc

    doc = Document(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text)
    for table_index, table in enumerate(doc.tables, 1):
        parts.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip() + "\n", "python-docx"


def materialize_text(path: Path) -> tuple[str, str]:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return path.read_text(encoding="utf-8", errors="replace"), "utf8-text"
    if suffix in HTML_EXTENSIONS:
        parser = _HTMLTextExtractor()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return parser.text(), "stdlib-html-parser"
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    raise RuntimeError(
        f"unsupported full-text format {suffix!r} for {path}; "
        "materialise it to UTF-8 text first"
    )


def resolve_slr_root(explicit: Path | None) -> Path:
    if explicit is not None:
        root = explicit.resolve()
    elif os.environ.get("SLR_REPO_ROOT"):
        root = Path(os.environ["SLR_REPO_ROOT"]).resolve()
    else:
        root = (Path(__file__).resolve().parents[3] / "slr").resolve()
    if not (root / "Cargo.toml").exists():
        raise SystemExit(f"SLR repository root not found: {root}")
    return root


def choose_ledger(artifact_root: Path, explicit: Path | None) -> Path:
    if explicit is not None:
        return explicit.resolve()
    reviewed = artifact_root / "screening_ledger_reviewed.tsv"
    unresolved = artifact_root / "screening_ledger.tsv"
    if reviewed.exists():
        return reviewed
    if unresolved.exists():
        return unresolved
    raise FileNotFoundError(
        f"no authoritative screening ledger under {artifact_root}"
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--slr-root", type=Path)
    ap.add_argument(
        "--artifact-root",
        type=Path,
        default=Path("artifacts/digital-esd/real-eric"),
    )
    ap.add_argument("--fulltext-index", type=Path)
    ap.add_argument("--screening-ledger", type=Path)
    ap.add_argument("--output-dir", type=Path)
    ap.add_argument(
        "--max-items",
        type=int,
        help="parse at most N verified artifacts; omitted means all verified artifacts",
    )
    ap.add_argument("--allow-partial", action="store_true")
    ap.add_argument(
        "--materialize-only",
        action="store_true",
        help="stop after verified bytes -> materialised UTF-8 + processing ledger; do not run the legacy scholarly parser",
    )
    ap.add_argument("--slr-review-receipts", type=Path)
    ap.add_argument("--source-audit-receipts", type=Path)
    args = ap.parse_args()

    slr_root = resolve_slr_root(args.slr_root)
    artifact_root = args.artifact_root
    if not artifact_root.is_absolute():
        artifact_root = (slr_root / artifact_root).resolve()
    else:
        artifact_root = artifact_root.resolve()

    if args.fulltext_index:
        fulltext_index = args.fulltext_index
        if not fulltext_index.is_absolute():
            fulltext_index = (slr_root / fulltext_index).resolve()
        else:
            fulltext_index = fulltext_index.resolve()
    else:
        fulltext_index = artifact_root / "fulltext" / "digital_esd_fulltext_index.tsv"
    ledger = choose_ledger(artifact_root, args.screening_ledger)
    output_dir = (
        args.output_dir.resolve()
        if args.output_dir
        else artifact_root / "slr-parse"
    )
    output_dir.mkdir(parents=True, exist_ok=True)

    if not fulltext_index.exists():
        raise FileNotFoundError(fulltext_index)

    rows = read_tsv(fulltext_index)
    verified_all = [row for row in rows if row.get("status") == "verified"]
    if not verified_all:
        raise RuntimeError("full-text index contains no verified artifacts")

    verified_all.sort(key=lambda row: str(row.get("source_identity_reference") or ""))
    if args.max_items is not None:
        if args.max_items < 1:
            raise ValueError("--max-items must be >= 1")
        verified = verified_all[: args.max_items]
    else:
        verified = verified_all

    materialized_dir = output_dir / "materialized-text"
    materialized_dir.mkdir(parents=True, exist_ok=True)

    materialization_receipts: list[dict[str, Any]] = []
    parser_inputs: list[dict[str, Any]] = []
    handoff_receipts: list[dict[str, Any]] = []
    failures: list[dict[str, Any]] = []

    for row in verified:
        ref = str(row["source_identity_reference"])
        artifact = Path(str(row.get("artifact_path") or ""))
        if not artifact.is_absolute():
            candidates = [
                (slr_root / artifact).resolve(),
                (artifact_root / artifact).resolve(),
                artifact.resolve(),
            ]
            artifact = next((p for p in candidates if p.exists()), candidates[0])
        else:
            artifact = artifact.resolve()

        original_digest = str(
            row.get("observed_sha256") or row.get("expected_sha256") or ""
        ).lower().removeprefix("sha256:")
        if not artifact.exists():
            failures.append({"source_identity_reference": ref, "reason": "artifact-missing"})
            continue
        observed = sha256_file(artifact)
        if observed != original_digest:
            failures.append({
                "source_identity_reference": ref,
                "reason": "artifact-digest-mismatch",
                "expected": original_digest,
                "observed": observed,
            })
            continue

        try:
            text, backend = materialize_text(artifact)
        except Exception as exc:
            failures.append({
                "source_identity_reference": ref,
                "reason": "text-materialization-failed",
                "error": str(exc),
                "artifact_path": str(artifact),
            })
            continue

        text_bytes = text.encode("utf-8")
        text_digest = sha256_bytes(text_bytes)
        materialized = materialized_dir / f"{text_digest}.txt"
        if not materialized.exists():
            materialized.write_bytes(text_bytes)

        original_revision = f"artifact-sha256:{original_digest}"
        parser_revision = f"materialized-text-sha256:{text_digest}"
        receipt_ref = f"materialization:{ref}:{text_digest}"

        materialization_receipts.append({
            "schema": "digital-esd-scholarly-text-materialization-v1",
            "materialization_receipt_reference": receipt_ref,
            "source_identity_reference": ref,
            "parent_source_revision_reference": original_revision,
            "parent_content_sha256": original_digest,
            "parent_artifact_path": str(artifact),
            "source_revision_reference": parser_revision,
            "content_sha256": text_digest,
            "artifact_path": str(materialized),
            "extraction_backend": backend,
            "candidate_only": True,
            "creates_semantic_authority": False,
            "creates_source_audit_admission": False,
        })

        parser_inputs.append({
            "source_identity_reference": ref,
            "source_revision_reference": parser_revision,
            "content_sha256": text_digest,
            "artifact_path": str(materialized),
            "parent_source_revision_reference": original_revision,
            "parent_content_sha256": original_digest,
            "materialization_receipt_reference": receipt_ref,
            "candidate_only": True,
            "creates_semantic_authority": False,
            "applicability_promoted": False,
            "claim_truth_promoted": False,
        })

        handoff_receipts.append({
            "schema": "digital-esd-slr-handoff-v1",
            "source_identity_reference": ref,
            "source_revision_reference": parser_revision,
            "parent_source_revision_reference": original_revision,
            "materialization_receipt_reference": receipt_ref,
            "handoff_status": "handed-to-slr",
            "handed_to_slr": True,
            "candidate_only": True,
            "creates_source_truth": False,
            "creates_source_audit_admission": False,
        })

    if failures and not args.allow_partial:
        failure_path = output_dir / "materialization-failures.jsonl"
        write_jsonl(failure_path, failures)
        raise RuntimeError(
            f"{len(failures)} verified artifacts failed text materialisation; "
            f"see {failure_path}; use --allow-partial only if partiality is intended"
        )

    materialization_path = output_dir / "materialization-receipts.jsonl"
    parser_input_path = output_dir / "scholarly-parser-input.jsonl"
    handoff_path = output_dir / "slr-handoff-receipts.jsonl"
    write_jsonl(materialization_path, materialization_receipts)
    write_jsonl(parser_input_path, parser_inputs)
    if failures:
        write_jsonl(output_dir / "materialization-failures.jsonl", failures)

    if args.materialize_only:
        processing_ledger_builder = HERE / "build_processing_ledger.py"
        processing_ledger = output_dir / "study-processing-ledger.jsonl"
        processing_manifest = output_dir / "study-processing-ledger-manifest.json"
        processing_cmd = [
            sys.executable,
            str(processing_ledger_builder),
            "--screening-ledger",
            str(ledger),
            "--fulltext-index",
            str(fulltext_index),
            "--materialization-receipts",
            str(materialization_path),
            "--output-ledger",
            str(processing_ledger),
            "--output-manifest",
            str(processing_manifest),
        ]
        existing_handoff = output_dir / "slr-handoff-receipts.jsonl"
        existing_parse = output_dir / "slr-parse-receipts.jsonl"
        existing_review = (
            args.slr_review_receipts.resolve()
            if args.slr_review_receipts
            else output_dir / "slr-review-receipts.jsonl"
        )
        existing_audit = (
            args.source_audit_receipts.resolve()
            if args.source_audit_receipts
            else output_dir / "source-audit-receipts.jsonl"
        )
        for flag, path in [
            ("--slr-handoff", existing_handoff),
            ("--slr-parse-receipts", existing_parse),
            ("--slr-review-receipts", existing_review),
            ("--source-audit-receipts", existing_audit),
        ]:
            if path.exists():
                processing_cmd.extend([flag, str(path)])
        print("+", " ".join(processing_cmd), file=sys.stderr)
        subprocess.run(
            processing_cmd,
            cwd=Path(__file__).resolve().parents[2],
            check=True,
        )

        receipt = {
            "schema": "digital-esd-verified-fulltext-materialization-run-v1",
            "verified_fulltext_total_count": len(verified_all),
            "selected_for_materialization_count": len(verified),
            "materialized_text_count": len(materialization_receipts),
            "materialization_failure_count": len(failures),
            "materialization_receipts_reference": str(materialization_path),
            "materialization_receipts_sha256": sha256_file(materialization_path),
            "study_processing_ledger_reference": str(processing_ledger),
            "study_processing_ledger_sha256": sha256_file(processing_ledger),
            "legacy_scholarly_parser_invoked": False,
            "materialization_creates_reviewed_evidence": False,
            "materialization_creates_source_audit_admission": False,
        }
        receipt_path = output_dir / "verified-fulltext-materialization-run.json"
        receipt_path.write_text(
            json.dumps(receipt, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        print(json.dumps(receipt, indent=2, sort_keys=True))
        return 0

    write_jsonl(handoff_path, handoff_receipts)

    scholarly = slr_root / "interop_scripts" / "digital_esd" / "scholarly_fulltext.py"
    parser_output_dir = output_dir / "parser"
    cmd = [
        sys.executable,
        str(scholarly),
        "run",
        "--input",
        str(parser_input_path),
        "--output-dir",
        str(parser_output_dir),
    ]
    if args.allow_partial:
        cmd.append("--allow-partial")
    print("+", " ".join(cmd), file=sys.stderr)
    env = os.environ.copy()
    existing_pythonpath = env.get("PYTHONPATH", "")
    env["PYTHONPATH"] = (
        str(slr_root)
        if not existing_pythonpath
        else str(slr_root) + os.pathsep + existing_pythonpath
    )
    subprocess.run(cmd, cwd=slr_root, env=env, check=True)

    verified_parse_path = parser_output_dir / "verified.jsonl"
    parsed = read_jsonl(verified_parse_path)
    if parser_inputs and not parsed and not args.allow_partial:
        raise RuntimeError(
            "generic SLR scholarly parser produced zero verified parse rows for "
            f"{len(parser_inputs)} prepared inputs; refusing silent ImportError/empty parse"
        )
    parsed_by_ref = {
        str(row.get("source_identity_reference") or ""): row for row in parsed
    }

    parse_receipts: list[dict[str, Any]] = []
    for handoff in handoff_receipts:
        ref = str(handoff["source_identity_reference"])
        parsed_row = parsed_by_ref.get(ref)
        if parsed_row is None:
            continue
        parse_receipts.append({
            "schema": "digital-esd-slr-parse-receipt-v1",
            "source_identity_reference": ref,
            "source_revision_reference": parsed_row.get("source_revision_reference"),
            "content_sha256": parsed_row.get("content_sha256"),
            "document_node_count": parsed_row.get("document_node_count", 0),
            "study_facet_count": parsed_row.get("study_facet_count", 0),
            "parsed": True,
            "parse_success": True,
            "candidate_only": True,
            "creates_semantic_authority": False,
            "applicability_promoted": False,
            "claim_truth_promoted": False,
            "creates_source_audit_admission": False,
        })

    parse_receipts_path = output_dir / "slr-parse-receipts.jsonl"
    write_jsonl(parse_receipts_path, parse_receipts)

    if len(parse_receipts) != len(parser_inputs) and not args.allow_partial:
        raise RuntimeError(
            "SLR parser denominator mismatch: "
            f"prepared={len(parser_inputs)} parsed={len(parse_receipts)}"
        )

    census = slr_root / "scripts" / "census_digital_esd_study_processing.py"
    census_output = output_dir / "study_processing_census_with_parse.json"
    census_cmd = [
        sys.executable,
        str(census),
        "--screening-ledger",
        str(ledger),
        "--fulltext-index",
        str(fulltext_index),
        "--slr-handoff",
        str(handoff_path),
        "--slr-parse-receipts",
        str(parse_receipts_path),
        "--output",
        str(census_output),
    ]
    if args.slr_review_receipts:
        census_cmd.extend(["--slr-review-receipts", str(args.slr_review_receipts.resolve())])
    if args.source_audit_receipts:
        census_cmd.extend(["--source-audit-receipts", str(args.source_audit_receipts.resolve())])
    print("+", " ".join(census_cmd), file=sys.stderr)
    subprocess.run(census_cmd, cwd=slr_root, check=True)

    processing_ledger_builder = HERE / "build_processing_ledger.py"
    processing_ledger = output_dir / "study-processing-ledger.jsonl"
    processing_manifest = output_dir / "study-processing-ledger-manifest.json"
    processing_cmd = [
        sys.executable,
        str(processing_ledger_builder),
        "--screening-ledger",
        str(ledger),
        "--fulltext-index",
        str(fulltext_index),
        "--materialization-receipts",
        str(materialization_path),
        "--slr-handoff",
        str(handoff_path),
        "--slr-parse-receipts",
        str(parse_receipts_path),
        "--output-ledger",
        str(processing_ledger),
        "--output-manifest",
        str(processing_manifest),
    ]
    if args.slr_review_receipts:
        processing_cmd.extend([
            "--slr-review-receipts",
            str(args.slr_review_receipts.resolve()),
        ])
    if args.source_audit_receipts:
        processing_cmd.extend([
            "--source-audit-receipts",
            str(args.source_audit_receipts.resolve()),
        ])
    print("+", " ".join(processing_cmd), file=sys.stderr)
    subprocess.run(processing_cmd, cwd=Path(__file__).resolve().parents[2], check=True)

    retrieval_residual_builder = HERE / "build_fulltext_retrieval_residual.py"
    retrieval_residual = output_dir / "fulltext-retrieval-residual.jsonl"
    retrieval_residual_manifest = output_dir / "fulltext-retrieval-residual-manifest.json"
    retrieval_cmd = [
        sys.executable,
        str(retrieval_residual_builder),
        "--processing-ledger",
        str(processing_ledger),
        "--output",
        str(retrieval_residual),
        "--manifest",
        str(retrieval_residual_manifest),
    ]
    print("+", " ".join(retrieval_cmd), file=sys.stderr)
    subprocess.run(retrieval_cmd, cwd=Path(__file__).resolve().parents[2], check=True)

    receipt = {
        "schema": "digital-esd-verified-fulltext-parse-run-v1",
        "verified_fulltext_total_count": len(verified_all),
        "selected_for_parse_count": len(verified),
        "materialized_text_count": len(materialization_receipts),
        "materialization_failure_count": len(failures),
        "handed_to_slr_count": len(handoff_receipts),
        "successfully_parsed_by_slr_count": len(parse_receipts),
        "fulltext_index_reference": str(fulltext_index),
        "fulltext_index_sha256": sha256_file(fulltext_index),
        "materialization_receipts_reference": str(materialization_path),
        "materialization_receipts_sha256": sha256_file(materialization_path),
        "slr_handoff_reference": str(handoff_path),
        "slr_handoff_sha256": sha256_file(handoff_path),
        "slr_parse_receipts_reference": str(parse_receipts_path),
        "slr_parse_receipts_sha256": sha256_file(parse_receipts_path),
        "census_reference": str(census_output),
        "census_sha256": sha256_file(census_output),
        "study_processing_ledger_reference": str(processing_ledger),
        "study_processing_ledger_sha256": sha256_file(processing_ledger),
        "study_processing_manifest_reference": str(processing_manifest),
        "study_processing_manifest_sha256": sha256_file(processing_manifest),
        "fulltext_retrieval_residual_reference": str(retrieval_residual),
        "fulltext_retrieval_residual_sha256": sha256_file(retrieval_residual),
        "fulltext_retrieval_residual_manifest_reference": str(retrieval_residual_manifest),
        "fulltext_retrieval_residual_manifest_sha256": sha256_file(retrieval_residual_manifest),
        "verified_bytes_count_as_parsed": False,
        "parse_creates_reviewed_evidence": False,
        "parse_creates_source_audit_admission": False,
    }
    receipt_path = output_dir / "verified-fulltext-parse-run.json"
    receipt_path.write_text(
        json.dumps(receipt, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(receipt, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
